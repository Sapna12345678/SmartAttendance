# ============================================================
# SMART ATTENDANCE SYSTEM
# attendance/views.py
# PART 1 / 2
# ============================================================

import base64
import json
import math
from io import BytesIO
from urllib.parse import quote

import qrcode

from django.conf import settings
from django.contrib import messages
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.http import JsonResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import reverse
from django.utils import timezone
from django.views.decorators.csrf import csrf_exempt

from .models import (
    Student,
    Subject,
    AttendanceSession,
    Attendance,
)


# ============================================================
# PUBLIC HOST
# ============================================================

def get_public_host(request=None):
    """
    Returns the stable public host used in QR codes.
    """

    host = getattr(
        settings,
        "SMART_ATTENDANCE_HOST",
        "smartattendance-edd1.onrender.com",
    )

    host = str(host).strip()

    if host.startswith("http://"):
        return host

    if host.startswith("https://"):
        return host

    return "https://" + host


# ============================================================
# HOME
# ============================================================

def home(request):
    """
    Main landing page.
    """

    if request.user.is_authenticated:

        # Student
        try:
            request.user.student_profile
            return redirect("student_dashboard")
        except Student.DoesNotExist:
            pass

        # Django admin/staff
        if request.user.is_staff:
            return redirect("/admin/")

        # Teacher session
        if request.session.get("teacher_logged_in"):
            return redirect("teacher_dashboard")

    return render(
        request,
        "attendance/home.html"
    )


# ============================================================
# TEACHER REQUIRED
# ============================================================

def teacher_required(view_func):
    """
    Protect teacher dashboard pages.
    """

    def wrapper(request, *args, **kwargs):

        if not request.session.get("teacher_logged_in"):

            return redirect(
                "teacher_login"
            )

        return view_func(
            request,
            *args,
            **kwargs
        )

    wrapper.__name__ = view_func.__name__

    return wrapper


# ============================================================
# TEACHER LOGIN
# ============================================================

def teacher_login(request):

    if request.session.get("teacher_logged_in"):

        return redirect(
            "teacher_dashboard"
        )

    if request.method == "POST":

        username = request.POST.get(
            "username",
            ""
        ).strip()

        password = request.POST.get(
            "password",
            ""
        )

        user = authenticate(
            request,
            username=username,
            password=password
        )

        if user is not None:

            # Teacher login
            request.session[
                "teacher_logged_in"
            ] = True

            request.session[
                "teacher_user_id"
            ] = user.id

            request.session[
                "teacher_username"
            ] = user.username

            return redirect(
                "teacher_dashboard"
            )

        messages.error(
            request,
            "Invalid Teacher ID or Password."
        )

    return render(
        request,
        "attendance/teacher_login.html"
    )


# ============================================================
# TEACHER LOGOUT
# ============================================================

def teacher_logout(request):

    request.session.pop(
        "teacher_logged_in",
        None
    )

    request.session.pop(
        "teacher_user_id",
        None
    )

    request.session.pop(
        "teacher_username",
        None
    )

    return redirect(
        "teacher_login"
    )


# ============================================================
# TEACHER DASHBOARD
# ============================================================

@teacher_required
def teacher_dashboard(request):

    today = timezone.localdate()

    total_students = Student.objects.count()

    present_today = Attendance.objects.filter(
        date=today,
        status="Present"
    ).count()

    absent_today = Attendance.objects.filter(
        date=today,
        status="Absent"
    ).count()

    total_today = present_today + absent_today

    if total_today > 0:
        attendance_percentage = round(
            (present_today / total_today) * 100,
            1
        )
    else:
        attendance_percentage = 0

    active_sessions = AttendanceSession.objects.filter(
        active=True,
        expires_at__gt=timezone.now()
    ).select_related(
        "subject",
        "teacher"
    ).order_by(
        "-created_at"
    )

    active_qr_count = active_sessions.count()

    recent_attendance = Attendance.objects.select_related(
        "student",
        "subject"
    ).order_by(
        "-date",
        "-time"
    )[:10]

    context = {

        "total_students": total_students,

        "present_today": present_today,

        "absent_today": absent_today,

        "attendance_percentage":
            attendance_percentage,

        "active_sessions":
            active_sessions,

        "active_qr_count":
            active_qr_count,

        "recent_attendance":
            recent_attendance,
    }

    return render(
        request,
        "attendance/teacher_dashboard.html",
        context
    )


# ============================================================
# GENERATE QR
# ============================================================

@teacher_required
def generate_qr(request):

    # --------------------------------------------------------
    # GET without generate = show form only
    # --------------------------------------------------------

    if request.method == "GET" and request.GET.get(
        "generate"
    ) != "1":

        subjects = Subject.objects.filter(
            active=True
        ).order_by(
            "branch",
            "semester",
            "name"
        )

        return render(
            request,
            "attendance/generate_qr.html",
            {
                "subjects": subjects,
            }
        )

    # --------------------------------------------------------
    # Read form data
    # --------------------------------------------------------

    if request.method == "POST":

        subject_id = request.POST.get(
            "subject"
        )

        branch = request.POST.get(
            "branch",
            ""
        ).strip()

        section = request.POST.get(
            "section",
            ""
        ).strip()

        semester = request.POST.get(
            "semester",
            ""
        ).strip()

        duration = request.POST.get(
            "duration",
            "60"
        )

    else:

        subject_id = request.GET.get(
            "subject"
        )

        branch = request.GET.get(
            "branch",
            ""
        ).strip()

        section = request.GET.get(
            "section",
            ""
        ).strip()

        semester = request.GET.get(
            "semester",
            ""
        ).strip()

        duration = request.GET.get(
            "duration",
            "60"
        )

    # --------------------------------------------------------
    # Validate subject
    # --------------------------------------------------------

    try:

        subject = Subject.objects.get(
            id=subject_id,
            active=True
        )

    except (
        Subject.DoesNotExist,
        ValueError,
        TypeError
    ):

        messages.error(
            request,
            "Please select a valid subject."
        )

        return redirect(
            "generate_qr"
        )

    # --------------------------------------------------------
    # Validate branch
    # --------------------------------------------------------

    if not branch:

        messages.error(
            request,
            "Branch is required."
        )

        return redirect(
            "generate_qr"
        )

    # --------------------------------------------------------
    # Validate section
    # --------------------------------------------------------

    if not section:

        messages.error(
            request,
            "Section is required."
        )

        return redirect(
            "generate_qr"
        )

    # --------------------------------------------------------
    # Semester
    # --------------------------------------------------------

    try:

        if isinstance(
            semester,
            str
        ) and semester.lower().startswith(
            "semester"
        ):

            semester_number = int(
                semester.split()[-1]
            )

        else:

            semester_number = int(
                semester
            )

    except (
        ValueError,
        TypeError
    ):

        semester_number = 3

    # Keep semester within normal range
    semester_number = max(
        1,
        min(
            semester_number,
            8
        )
    )

    # --------------------------------------------------------
    # Duration
    # --------------------------------------------------------

    try:

        duration_seconds = int(
            duration
        )

    except (
        ValueError,
        TypeError
    ):

        duration_seconds = 60

    duration_seconds = max(
        30,
        min(
            duration_seconds,
            300
        )
    )

    # --------------------------------------------------------
    # Teacher
    # --------------------------------------------------------

    teacher = None

    teacher_user_id = request.session.get(
        "teacher_user_id"
    )

    if teacher_user_id:

        try:

            from django.contrib.auth.models import User

            teacher = User.objects.get(
                id=teacher_user_id
            )

        except User.DoesNotExist:

            teacher = None

    # --------------------------------------------------------
    # Deactivate previous QR sessions
    # --------------------------------------------------------

    AttendanceSession.objects.filter(
        active=True
    ).update(
        active=False
    )

    # --------------------------------------------------------
    # Create new attendance session
    # --------------------------------------------------------

    now = timezone.now()

    expires_at = (
        now +
        timezone.timedelta(
            seconds=duration_seconds
        )
    )

    session = AttendanceSession.objects.create(

        subject=subject,

        teacher=teacher,

        branch=branch,

        section=section,

        semester=semester_number,

        created_at=now,

        expires_at=expires_at,

        active=True,
    )

    # --------------------------------------------------------
    # Stable public scan URL
    # --------------------------------------------------------

    public_host = get_public_host(
        request
    )

    scan_url = (
        public_host +
        "/scan/?session=" +
        str(session.session_id)
    )

    # --------------------------------------------------------
    # Generate QR in memory
    # --------------------------------------------------------

    qr = qrcode.QRCode(

        version=1,

        error_correction=qrcode.constants.ERROR_CORRECT_H,

        box_size=10,

        border=4,
    )

    qr.add_data(
        scan_url
    )

    qr.make(
        fit=True
    )

    qr_image = qr.make_image(
        fill_color="black",
        back_color="white"
    )

    # --------------------------------------------------------
    # Convert QR to Base64
    # --------------------------------------------------------

    buffer = BytesIO()

    qr_image.save(
        buffer,
        format="PNG"
    )

    qr_base64 = base64.b64encode(
        buffer.getvalue()
    ).decode(
        "utf-8"
    )

    qr_image_base64 = (
        "data:image/png;base64," +
        qr_base64
    )

    # --------------------------------------------------------
    # Filename
    # --------------------------------------------------------

    filename = (
        "attendance_qr_" +
        str(session.session_id) +
        ".png"
    )

    # --------------------------------------------------------
    # Render
    # --------------------------------------------------------

    context = {

        "session":
            session,

        "session_id":
            str(session.session_id),

        "subject":
            subject,

        "branch":
            branch,

        "section":
            section,

        "semester":
            semester_number,

        "duration":
            duration_seconds,

        "remaining_seconds":
            duration_seconds,

        "qr_image_base64":
            qr_image_base64,

        "qr_image_url":
            qr_image_base64,

        "scan_url":
            scan_url,

        "filename":
            filename,
    }

    return render(
        request,
        "attendance/generate_qr.html",
        context
    )


# ============================================================
# PROJECTOR
# ============================================================

@teacher_required
def projector(request):

    session_id = request.GET.get(
        "session"
    )

    if not session_id:

        return redirect(
            "generate_qr"
        )

    try:

        session = AttendanceSession.objects.select_related(
            "subject"
        ).get(
            session_id=session_id
        )

    except AttendanceSession.DoesNotExist:

        messages.error(
            request,
            "Attendance session not found."
        )

        return redirect(
            "generate_qr"
        )

    # --------------------------------------------------------
    # Calculate remaining time
    # --------------------------------------------------------

    remaining_seconds = max(
        0,
        int(
            (
                session.expires_at -
                timezone.now()
            ).total_seconds()
        )
    )

    if remaining_seconds <= 0:

        session.active = False
        session.save(
            update_fields=["active"]
        )

    # --------------------------------------------------------
    # Generate QR again for projector
    # --------------------------------------------------------

    public_host = get_public_host(
        request
    )

    scan_url = (
        public_host +
        "/scan/?session=" +
        str(session.session_id)
    )

    qr = qrcode.QRCode(

        version=1,

        error_correction=qrcode.constants.ERROR_CORRECT_H,

        box_size=10,

        border=4,
    )

    qr.add_data(
        scan_url
    )

    qr.make(
        fit=True
    )

    qr_image = qr.make_image(
        fill_color="black",
        back_color="white"
    )

    buffer = BytesIO()

    qr_image.save(
        buffer,
        format="PNG"
    )

    qr_base64 = base64.b64encode(
        buffer.getvalue()
    ).decode(
        "utf-8"
    )

    qr_image_base64 = (
        "data:image/png;base64," +
        qr_base64
    )

    duration = max(
        1,
        int(
            (
                session.expires_at -
                session.created_at
            ).total_seconds()
        )
    )

    context = {

        "session":
            session,

        "session_id":
            str(session.session_id),

        "subject":
            session.subject,

        "branch":
            session.branch,

        "section":
            session.section,

        "semester":
            session.semester,

        "remaining_seconds":
            remaining_seconds,

        "duration":
            duration,

        "qr_image_base64":
            qr_image_base64,

        "scan_url":
            scan_url,
    }

    return render(
        request,
        "attendance/projector.html",
        context
    )


# ============================================================
# SCAN QR
# ============================================================

def scan_qr(request):

    session_id = request.GET.get(
        "session"
    )

    # --------------------------------------------------------
    # No session ID
    # --------------------------------------------------------

    if not session_id:

        return render(
            request,
            "attendance/scan.html",
            {
                "valid": False,
                "error":
                    "Attendance QR session is missing."
            }
        )

    # --------------------------------------------------------
    # Find session
    # --------------------------------------------------------

    try:

        session = AttendanceSession.objects.select_related(
            "subject"
        ).get(
            session_id=session_id
        )

    except (
        AttendanceSession.DoesNotExist,
        ValueError,
        TypeError
    ):

        return render(
            request,
            "attendance/scan.html",
            {
                "valid": False,
                "error":
                    "Invalid Attendance QR."
            }
        )

    # --------------------------------------------------------
    # Check expiry
    # --------------------------------------------------------

    now = timezone.now()

    if (
        not session.active
        or
        session.expires_at <= now
    ):

        if session.active:

            session.active = False

            session.save(
                update_fields=["active"]
            )

        return render(
            request,
            "attendance/scan.html",
            {
                "valid": False,
                "expired": True,
                "error":
                    "This attendance QR has expired."
            }
        )

    # --------------------------------------------------------
    # Remaining time
    # --------------------------------------------------------

    time_left = max(
        0,
        int(
            (
                session.expires_at -
                now
            ).total_seconds()
        )
    )

    # --------------------------------------------------------
    # Identify logged-in student
    # --------------------------------------------------------

    student = None

    if request.user.is_authenticated:

        try:

            student = request.user.student_profile

        except Student.DoesNotExist:

            student = None

    # --------------------------------------------------------
    # Student not logged in
    # --------------------------------------------------------

    if student is None:

        login_url = (
            reverse("student_login") +
            "?next=" +
            quote(
                request.get_full_path()
            )
        )

        return render(
            request,
            "attendance/scan.html",
            {
                "valid": True,

                "requires_login": True,

                "login_url": login_url,

                "session_id":
                    str(session.session_id),

                "subject":
                    session.subject,

                "branch":
                    session.branch,

                "section":
                    session.section,

                "semester":
                    session.semester,

                "time_left":
                    time_left,
            }
        )

    # --------------------------------------------------------
    # Logged-in student
    # --------------------------------------------------------

    return render(
        request,
        "attendance/scan.html",
        {
            "valid": True,

            "requires_login": False,

            "session_id":
                str(session.session_id),

            "subject":
                session.subject,

            "branch":
                session.branch,

            "section":
                session.section,

            "semester":
                session.semester,

            "time_left":
                time_left,

            "student":
                student,

            "student_name":
                student.name,

            "student_roll_no":
                student.roll_no,

            "student_course":
                student.course,

            "student_year":
                student.year,
        }
    )


# ============================================================
# MARK ATTENDANCE
# ============================================================

@csrf_exempt
def mark_attendance(request):

    if request.method != "POST":

        return JsonResponse(
            {
                "success": False,
                "error":
                    "POST request required."
            },
            status=405
        )

    # --------------------------------------------------------
    # Parse request data
    # --------------------------------------------------------

    data = {}

    try:

        if request.body:

            data = json.loads(
                request.body.decode(
                    "utf-8"
                )
            )

    except (
        json.JSONDecodeError,
        UnicodeDecodeError
    ):

        data = request.POST

    # --------------------------------------------------------
    # Session ID
    # --------------------------------------------------------

    session_id = data.get(
        "session_id"
    )

    if not session_id:

        return JsonResponse(
            {
                "success": False,
                "error":
                    "Attendance session is required."
            },
            status=400
        )

    # --------------------------------------------------------
    # Get session
    # --------------------------------------------------------

    try:

        session = AttendanceSession.objects.select_related(
            "subject"
        ).get(
            session_id=session_id
        )

    except (
        AttendanceSession.DoesNotExist,
        ValueError,
        TypeError
    ):

        return JsonResponse(
            {
                "success": False,
                "error":
                    "Invalid attendance session."
            },
            status=404
        )

    # --------------------------------------------------------
    # Check session expiry
    # --------------------------------------------------------

    now = timezone.now()

    if (
        not session.active
        or
        session.expires_at <= now
    ):

        if session.active:

            session.active = False

            session.save(
                update_fields=["active"]
            )

        return JsonResponse(
            {
                "success": False,
                "error":
                    "Attendance QR has expired."
            },
            status=400
        )

    # --------------------------------------------------------
    # IMPORTANT:
    # Student identity comes from logged-in account.
    # Never trust student_id sent by browser.
    # --------------------------------------------------------

    if not request.user.is_authenticated:

        return JsonResponse(
            {
                "success": False,
                "error":
                    "Please login as a student first.",
                "requires_login": True,
            },
            status=401
        )

    try:

        student = request.user.student_profile

    except Student.DoesNotExist:

        return JsonResponse(
            {
                "success": False,
                "error":
                    "Student profile not found."
            },
            status=403
        )

    # --------------------------------------------------------
    # Student eligibility
    # --------------------------------------------------------

    session_branch = (
        str(session.branch or "")
        .strip()
        .lower()
    )

    student_branch = (
        str(student.branch or "")
        .strip()
        .lower()
    )

    if (
        session_branch
        and
        student_branch
        and
        session_branch != student_branch
    ):

        return JsonResponse(
            {
                "success": False,

                "status":
                    "Absent",

                "error":
                    "Student branch does not match this attendance session.",

                "name":
                    student.name,

                "roll_no":
                    student.roll_no,
            },
            status=403
        )

    # --------------------------------------------------------
    # Section validation
    # --------------------------------------------------------

    session_section = (
        str(session.section or "")
        .strip()
        .lower()
    )

    student_section = (
        str(student.section or "")
        .strip()
        .lower()
    )

    if (
        session_section
        and
        student_section
        and
        session_section != student_section
    ):

        return JsonResponse(
            {
                "success": False,

                "status":
                    "Absent",

                "error":
                    "Student section does not match this attendance session.",

                "name":
                    student.name,

                "roll_no":
                    student.roll_no,
            },
            status=403
        )

    # --------------------------------------------------------
    # Semester validation
    # --------------------------------------------------------

    try:

        student_semester = int(
            student.semester
        ) if student.semester is not None else None

    except (
        ValueError,
        TypeError
    ):

        student_semester = None

    try:

        session_semester = int(
            session.semester
        ) if session.semester is not None else None

    except (
        ValueError,
        TypeError
    ):

        session_semester = None

    if (
        session_semester is not None
        and
        student_semester is not None
        and
        session_semester != student_semester
    ):

        return JsonResponse(
            {
                "success": False,

                "status":
                    "Absent",

                "error":
                    "Student semester does not match this attendance session.",

                "name":
                    student.name,

                "roll_no":
                    student.roll_no,
            },
            status=403
        )

    # --------------------------------------------------------
    # GPS coordinates
    # --------------------------------------------------------

    latitude = data.get(
        "latitude"
    )

    longitude = data.get(
        "longitude"
    )

    if latitude in (
        None,
        ""
    ) or longitude in (
        None,
        ""
    ):

        return JsonResponse(
            {
                "success": False,

                "error":
                    "Location permission is required. Please enable GPS."
            },
            status=400
        )

    # --------------------------------------------------------
    # Convert GPS values
    # --------------------------------------------------------

    try:

        latitude = float(
            latitude
        )

        longitude = float(
            longitude
        )

    except (
        ValueError,
        TypeError
    ):

        return JsonResponse(
            {
                "success": False,

                "error":
                    "Invalid GPS coordinates."
            },
            status=400
        )

    # --------------------------------------------------------
    # College location
    # UIM / Prayagraj demo location
    # --------------------------------------------------------

    college_lat = float(
        getattr(
            settings,
            "ATTENDANCE_LATITUDE",
            25.342787
        )
    )

    college_lon = float(
        getattr(
            settings,
            "ATTENDANCE_LONGITUDE",
            81.902116
        )
    )

    allowed_radius = float(
        getattr(
            settings,
            "ATTENDANCE_RADIUS_METERS",
            300
        )
    )

    # --------------------------------------------------------
    # Haversine distance
    # --------------------------------------------------------

    earth_radius = 6371000

    lat1 = math.radians(
        college_lat
    )

    lat2 = math.radians(
        latitude
    )

    delta_lat = math.radians(
        latitude - college_lat
    )

    delta_lon = math.radians(
        longitude - college_lon
    )

    a = (
        math.sin(
            delta_lat / 2
        ) ** 2
        +
        math.cos(lat1)
        *
        math.cos(lat2)
        *
        math.sin(
            delta_lon / 2
        ) ** 2
    )

    c = 2 * math.atan2(
        math.sqrt(a),
        math.sqrt(
            1 - a
        )
    )

    distance = (
        earth_radius * c
    )

    # --------------------------------------------------------
    # Location verification
    # --------------------------------------------------------

    location_verified = (
        distance <= allowed_radius
    )

    status = (
        "Present"
        if location_verified
        else "Absent"
    )

    # --------------------------------------------------------
    # Prevent duplicate attendance
    # --------------------------------------------------------

    today = timezone.localdate()

    existing = Attendance.objects.filter(
        student=student,
        session_id=session.session_id,
        date=today
    ).first()

    if existing:

        return JsonResponse(
            {
                "success": True,

                "already_marked": True,

                "status":
                    existing.status,

                "name":
                    student.name,

                "roll_no":
                    student.roll_no,

                "course":
                    student.course,

                "year":
                    student.year,

                "branch":
                    student.branch,

                "section":
                    student.section,

                "semester":
                    student.semester,

                "subject":
                    str(session.subject)
                    if session.subject
                    else "General",

                "distance":
                    round(
                        distance,
                        2
                    ),

                "allowed_radius":
                    allowed_radius,

                "location_verified":
                    existing.location_verified,
            }
        )

    # --------------------------------------------------------
    # Create attendance
    # --------------------------------------------------------

    attendance = Attendance.objects.create(

        student=student,

        session_id=session.session_id,

        subject=session.subject,

        date=today,

        time=timezone.localtime().time(),

        status=status,

        latitude=latitude,

        longitude=longitude,

        location_verified=location_verified,
    )

    # --------------------------------------------------------
    # Response
    # --------------------------------------------------------

    return JsonResponse(
        {
            "success": True,

            "already_marked": False,

            "status":
                attendance.status,

            "name":
                student.name,

            "roll_no":
                student.roll_no,

            "course":
                student.course,

            "year":
                student.year,

            "branch":
                student.branch,

            "section":
                student.section,

            "semester":
                student.semester,

            "subject":
                str(session.subject)
                if session.subject
                else "General",

            "distance":
                round(
                    distance,
                    2
                ),

            "allowed_radius":
                allowed_radius,

            "location_verified":
                location_verified,
        }
    )


# ============================================================
# TODAY ATTENDANCE
# ============================================================

@teacher_required
def today_attendance(request):

    today = timezone.localdate()

    attendances = Attendance.objects.select_related(
        "student",
        "subject"
    ).filter(
        date=today
    ).order_by(
        "-time"
    )

    total_records = attendances.count()

    present_count = attendances.filter(
        status="Present"
    ).count()

    absent_count = attendances.filter(
        status="Absent"
    ).count()

    if total_records:

        attendance_percentage = round(
            (
                present_count /
                total_records
            ) * 100,
            1
        )

    else:

        attendance_percentage = 0

    context = {

        "attendances":
            attendances,

        "records":
            attendances,

        "total_records":
            total_records,

        "present_count":
            present_count,

        "absent_count":
            absent_count,

        "attendance_percentage":
            attendance_percentage,

        "today":
            today,
    }

    return render(
        request,
        "attendance/today_attendance.html",
        context
    )
# ============================================================
# SMART ATTENDANCE SYSTEM
# attendance/views.py
# PART 2 / 2
# ============================================================


# ============================================================
# STUDENTS
# ============================================================

@teacher_required
def students(request):

    student_list = Student.objects.all().order_by(
        "roll_no"
    )

    context = {
        "students": student_list,
    }

    return render(
        request,
        "attendance/students.html",
        context
    )


# ============================================================
# ATTENDANCE LIST
# ============================================================

@teacher_required
def attendance_list(request):

    attendances = Attendance.objects.select_related(
        "student",
        "subject"
    ).all().order_by(
        "-date",
        "-time"
    )

    # --------------------------------------------------------
    # Date filter
    # --------------------------------------------------------

    selected_date = request.GET.get(
        "date",
        ""
    ).strip()

    if selected_date:

        try:

            attendances = attendances.filter(
                date=selected_date
            )

        except Exception:

            pass

    # --------------------------------------------------------
    # Search
    # --------------------------------------------------------

    search = request.GET.get(
        "search",
        ""
    ).strip()

    if search:

        from django.db.models import Q

        attendances = attendances.filter(

            Q(
                student__name__icontains=search
            )
            |
            Q(
                student__roll_no__icontains=search
            )
            |
            Q(
                subject__name__icontains=search
            )
            |
            Q(
                subject__code__icontains=search
            )
        )

    # --------------------------------------------------------
    # Statistics
    # --------------------------------------------------------

    total_records = attendances.count()

    present_count = attendances.filter(
        status="Present"
    ).count()

    absent_count = attendances.filter(
        status="Absent"
    ).count()

    if total_records:

        attendance_percentage = round(
            (
                present_count /
                total_records
            ) * 100,
            1
        )

    else:

        attendance_percentage = 0

    context = {

        "attendances":
            attendances,

        "records":
            attendances,

        "total_records":
            total_records,

        "present_count":
            present_count,

        "absent_count":
            absent_count,

        "attendance_percentage":
            attendance_percentage,

        "selected_date":
            selected_date,

        "search":
            search,
    }

    return render(
        request,
        "attendance/attendance_list.html",
        context
    )


# ============================================================
# STUDENT HISTORY
# ============================================================

@teacher_required
def student_history(request):

    student_id = request.GET.get(
        "student"
    )

    selected_student = None

    attendances = Attendance.objects.none()

    if student_id:

        try:

            selected_student = Student.objects.get(
                id=student_id
            )

            attendances = Attendance.objects.select_related(
                "student",
                "subject"
            ).filter(
                student=selected_student
            ).order_by(
                "-date",
                "-time"
            )

        except (
            Student.DoesNotExist,
            ValueError,
            TypeError
        ):

            selected_student = None

    # --------------------------------------------------------
    # Statistics
    # --------------------------------------------------------

    total_records = attendances.count()

    present_count = attendances.filter(
        status="Present"
    ).count()

    absent_count = attendances.filter(
        status="Absent"
    ).count()

    if total_records:

        attendance_percentage = round(
            (
                present_count /
                total_records
            ) * 100,
            1
        )

    else:

        attendance_percentage = 0

    student_list = Student.objects.all().order_by(
        "roll_no"
    )

    context = {

        "student":
            selected_student,

        "selected_student":
            selected_student,

        "students":
            student_list,

        "attendances":
            attendances,

        "records":
            attendances,

        "total_records":
            total_records,

        "present_count":
            present_count,

        "absent_count":
            absent_count,

        "attendance_percentage":
            attendance_percentage,
    }

    return render(
        request,
        "attendance/student_history.html",
        context
    )


# ============================================================
# SUBJECT-WISE ATTENDANCE
# ============================================================

@teacher_required
def subject_wise_attendance(request):

    subjects = Subject.objects.filter(
        active=True
    ).order_by(
        "branch",
        "semester",
        "name"
    )

    subject_id = request.GET.get(
        "subject"
    )

    selected_subject = None

    records = []

    if subject_id:

        try:

            selected_subject = Subject.objects.get(
                id=subject_id
            )

        except (
            Subject.DoesNotExist,
            ValueError,
            TypeError
        ):

            selected_subject = None

    if selected_subject:

        students = Student.objects.all().order_by(
            "roll_no"
        )

        for student in students:

            student_records = Attendance.objects.filter(
                student=student,
                subject=selected_subject
            )

            total = student_records.count()

            present = student_records.filter(
                status="Present"
            ).count()

            absent = student_records.filter(
                status="Absent"
            ).count()

            if total:

                percentage = round(
                    (
                        present /
                        total
                    ) * 100,
                    1
                )

            else:

                percentage = 0

            records.append(
                {
                    "student":
                        student,

                    "total":
                        total,

                    "present":
                        present,

                    "absent":
                        absent,

                    "percentage":
                        percentage,
                }
            )

    context = {

        "subjects":
            subjects,

        "selected_subject":
            selected_subject,

        "records":
            records,
    }

    return render(
        request,
        "attendance/subject_wise.html",
        context
    )


# ============================================================
# REPORTS
# ============================================================

@teacher_required
def reports(request):

    today = timezone.localdate()

    total_students = Student.objects.count()

    total_records = Attendance.objects.count()

    present_count = Attendance.objects.filter(
        status="Present"
    ).count()

    absent_count = Attendance.objects.filter(
        status="Absent"
    ).count()

    today_records = Attendance.objects.filter(
        date=today
    ).count()

    today_present = Attendance.objects.filter(
        date=today,
        status="Present"
    ).count()

    today_absent = Attendance.objects.filter(
        date=today,
        status="Absent"
    ).count()

    if total_records:

        overall_percentage = round(
            (
                present_count /
                total_records
            ) * 100,
            1
        )

    else:

        overall_percentage = 0

    if today_records:

        today_percentage = round(
            (
                today_present /
                today_records
            ) * 100,
            1
        )

    else:

        today_percentage = 0

    context = {

        "total_students":
            total_students,

        "total_records":
            total_records,

        "present_count":
            present_count,

        "absent_count":
            absent_count,

        "overall_percentage":
            overall_percentage,

        "today_records":
            today_records,

        "today_present":
            today_present,

        "today_absent":
            today_absent,

        "today_percentage":
            today_percentage,

        "today":
            today,
    }

    return render(
        request,
        "attendance/reports.html",
        context
    )


# ============================================================
# EXPORT ATTENDANCE TO EXCEL
# ============================================================

@teacher_required
def export_attendance(request):

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from django.http import HttpResponse

    attendances = Attendance.objects.select_related(
        "student",
        "subject"
    ).all().order_by(
        "-date",
        "-time"
    )

    # --------------------------------------------------------
    # Optional date filter
    # --------------------------------------------------------

    selected_date = request.GET.get(
        "date",
        ""
    ).strip()

    if selected_date:

        try:

            attendances = attendances.filter(
                date=selected_date
            )

        except Exception:

            pass

    # --------------------------------------------------------
    # Create workbook
    # --------------------------------------------------------

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = "Attendance Report"

    headers = [

        "Student Name",

        "Roll Number",

        "Course",

        "Year",

        "Branch",

        "Section",

        "Semester",

        "Subject",

        "Subject Code",

        "Date",

        "Time",

        "Status",

        "Latitude",

        "Longitude",

        "Location Verified",
    ]

    worksheet.append(
        headers
    )

    # --------------------------------------------------------
    # Header styling
    # --------------------------------------------------------

    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="173B6C"
        )

        cell.alignment = Alignment(
            horizontal="center"
        )

    # --------------------------------------------------------
    # Data
    # --------------------------------------------------------

    for attendance in attendances:

        student = attendance.student

        subject = attendance.subject

        worksheet.append(
            [

                student.name
                if student
                else "",

                student.roll_no
                if student
                else "",

                student.course
                if student
                else "",

                student.year
                if student
                else "",

                student.branch
                if student
                else "",

                student.section
                if student
                else "",

                student.semester
                if student
                else "",

                subject.name
                if subject
                else "General",

                subject.code
                if subject
                else "",

                attendance.date,

                attendance.time,

                attendance.status,

                attendance.latitude,

                attendance.longitude,

                "Yes"
                if attendance.location_verified
                else "No",
            ]
        )

    # --------------------------------------------------------
    # Column widths
    # --------------------------------------------------------

    widths = {

        "A": 24,
        "B": 15,
        "C": 15,
        "D": 10,
        "E": 12,
        "F": 12,
        "G": 12,
        "H": 28,
        "I": 16,
        "J": 14,
        "K": 14,
        "L": 14,
        "M": 15,
        "N": 15,
        "O": 18,
    }

    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width

    # --------------------------------------------------------
    # Response
    # --------------------------------------------------------

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )

    response[
        "Content-Disposition"
    ] = (
        'attachment; '
        'filename="smart_attendance_report.xlsx"'
    )

    workbook.save(
        response
    )

    return response


# ============================================================
# MARK ABSENT STUDENTS
# ============================================================

@teacher_required
def mark_absent(request):

    today = timezone.localdate()

    # --------------------------------------------------------
    # Active/recent session
    # --------------------------------------------------------

    session_id = request.GET.get(
        "session"
    )

    session = None

    if session_id:

        try:

            session = AttendanceSession.objects.get(
                session_id=session_id
            )

        except (
            AttendanceSession.DoesNotExist,
            ValueError,
            TypeError
        ):

            session = None

    # --------------------------------------------------------
    # If no session selected, find today's latest session
    # --------------------------------------------------------

    if session is None:

        session = AttendanceSession.objects.filter(
            created_at__date=today
        ).order_by(
            "-created_at"
        ).first()

    if session is None:

        messages.warning(
            request,
            "No attendance session found for today."
        )

        return redirect(
            "teacher_dashboard"
        )

    # --------------------------------------------------------
    # Eligible students
    # --------------------------------------------------------

    students = Student.objects.all()

    session_branch = (
        str(session.branch or "")
        .strip()
        .lower()
    )

    session_section = (
        str(session.section or "")
        .strip()
        .lower()
    )

    eligible_students = []

    for student in students:

        student_branch = (
            str(student.branch or "")
            .strip()
            .lower()
        )

        student_section = (
            str(student.section or "")
            .strip()
            .lower()
        )

        if (
            session_branch
            and student_branch
            and session_branch != student_branch
        ):
            continue

        if (
            session_section
            and student_section
            and session_section != student_section
        ):
            continue

        already_marked = Attendance.objects.filter(
            student=student,
            session_id=session.session_id,
            date=today
        ).exists()

        if not already_marked:

            eligible_students.append(
                student
            )

    # --------------------------------------------------------
    # Create Absent records
    # --------------------------------------------------------

    created_count = 0

    for student in eligible_students:

        Attendance.objects.create(

            student=student,

            session_id=session.session_id,

            subject=session.subject,

            date=today,

            time=timezone.localtime().time(),

            status="Absent",

            latitude=None,

            longitude=None,

            location_verified=False,
        )

        created_count += 1

    messages.success(
        request,
        f"{created_count} student(s) marked Absent."
    )

    return redirect(
        "today_attendance"
    )


# ============================================================
# CLEAN EXPIRED SESSIONS
# ============================================================

@teacher_required
def clean_expired_sessions(request):

    now = timezone.now()

    expired_sessions = AttendanceSession.objects.filter(
        active=True,
        expires_at__lte=now
    )

    count = expired_sessions.count()

    expired_sessions.update(
        active=False
    )

    messages.success(
        request,
        f"{count} expired QR session(s) cleaned."
    )

    return redirect(
        "teacher_dashboard"
    )


# ============================================================
# TEACHER PROFILE
# ============================================================

@teacher_required
def profile(request):

    from django.contrib.auth.models import User

    teacher = None

    teacher_user_id = request.session.get(
        "teacher_user_id"
    )

    if teacher_user_id:

        try:

            teacher = User.objects.get(
                id=teacher_user_id
            )

        except User.DoesNotExist:

            teacher = None

    if request.method == "POST" and teacher:

        email = request.POST.get(
            "email",
            ""
        ).strip()

        first_name = request.POST.get(
            "first_name",
            ""
        ).strip()

        last_name = request.POST.get(
            "last_name",
            ""
        ).strip()

        teacher.email = email

        teacher.first_name = first_name

        teacher.last_name = last_name

        teacher.save()

        messages.success(
            request,
            "Profile updated successfully."
        )

        return redirect(
            "profile"
        )

    context = {

        "teacher":
            teacher,
    }

    return render(
        request,
        "attendance/profile.html",
        context
    )


# ============================================================
# URL COMPATIBILITY FUNCTIONS
# ============================================================

@teacher_required
def students_page(request):

    student_list = Student.objects.all().order_by(
        "roll_no"
    )

    return render(
        request,
        "attendance/students.html",
        {
            "students":
                student_list
        }
    )


@teacher_required
def export_attendance_excel(request):

    return export_attendance(
        request
    )


@teacher_required
def mark_absent_students(request):

    return mark_absent(
        request
    )